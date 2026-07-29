from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def part_inventory(path: Path) -> dict[str, dict[str, object]]:
    out: dict[str, dict[str, object]] = {}
    with zipfile.ZipFile(path) as package:
        for info in package.infolist():
            data = package.read(info.filename)
            out[info.filename] = {
                "size": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            }
    return out


def style_record(style) -> dict[str, object]:
    pf = style.paragraph_format
    font = style.font
    return {
        "name": style.name,
        "type": str(style.type),
        "base_style": style.base_style.name if style.base_style else None,
        "font_name": font.name,
        "font_size_pt": font.size.pt if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
        "color": str(font.color.rgb) if font.color and font.color.rgb else None,
        "alignment": str(pf.alignment) if pf.alignment is not None else None,
        "space_before_pt": pf.space_before.pt if pf.space_before else None,
        "space_after_pt": pf.space_after.pt if pf.space_after else None,
        "line_spacing": (
            float(pf.line_spacing)
            if isinstance(pf.line_spacing, (int, float))
            else (pf.line_spacing.pt if pf.line_spacing else None)
        ),
        "keep_with_next": pf.keep_with_next,
        "page_break_before": pf.page_break_before,
    }


def cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def document_record(path: Path) -> dict[str, object]:
    doc = Document(path)
    paragraphs = []
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "index": i,
                    "style": paragraph.style.name if paragraph.style else None,
                    "text": text,
                }
            )
    tables = []
    for ti, table in enumerate(doc.tables):
        tables.append(
            {
                "index": ti,
                "style": table.style.name if table.style else None,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "first_row": [cell_text(c) for c in table.rows[0].cells] if table.rows else [],
            }
        )
    sections = []
    for i, section in enumerate(doc.sections):
        sections.append(
            {
                "index": i,
                "page_width_in": section.page_width.inches,
                "page_height_in": section.page_height.inches,
                "top_margin_in": section.top_margin.inches,
                "bottom_margin_in": section.bottom_margin.inches,
                "left_margin_in": section.left_margin.inches,
                "right_margin_in": section.right_margin.inches,
                "header_distance_in": section.header_distance.inches,
                "footer_distance_in": section.footer_distance.inches,
                "different_first_page": section.different_first_page_header_footer,
            }
        )

    full_text = "\n".join(p["text"] for p in paragraphs)
    abstract_match = re.search(r"\bAbstract\b(.*?)(?:\bKeywords?\b)", full_text, re.S | re.I)
    abstract_words = (
        len(re.findall(r"\b[\w'-]+\b", abstract_match.group(1))) if abstract_match else None
    )
    return {
        "path": str(path),
        "sha256": sha256(path),
        "size": path.stat().st_size,
        "sections": sections,
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "table_count": len(doc.tables),
        "tables": tables,
        "inline_shape_count": len(doc.inline_shapes),
        "abstract_words": abstract_words,
        "all_text_words": len(re.findall(r"\b[\w'-]+\b", full_text)),
        "core_properties": {
            "title": doc.core_properties.title,
            "subject": doc.core_properties.subject,
            "author": doc.core_properties.author,
            "keywords": doc.core_properties.keywords,
            "comments": doc.core_properties.comments,
            "last_modified_by": doc.core_properties.last_modified_by,
        },
        "styles": [
            style_record(style)
            for style in doc.styles
            if style.name
            in {
                "Normal",
                "Title",
                "Subtitle",
                "Heading 1",
                "Heading 2",
                "Heading 3",
                "Abstract",
                "Caption",
                "Image Caption",
                "Table Caption",
                "Bibliography",
                "Author",
                "Affiliation",
            }
        ],
    }


def main() -> None:
    if len(sys.argv) < 4:
        raise SystemExit("usage: inspect_tdp_docs.py TEMPLATE.DOTX REFERENCE.DOCX MANUSCRIPT.DOCX")
    template = Path(sys.argv[1])
    reference = Path(sys.argv[2])
    manuscript = Path(sys.argv[3])
    output = Path(sys.argv[4]) if len(sys.argv) > 4 else Path("inspection.json")

    template_parts = part_inventory(template)
    reference_parts = part_inventory(reference)
    manuscript_parts = part_inventory(manuscript)
    compared = {}
    for part in (
        "word/styles.xml",
        "word/theme/theme1.xml",
        "word/fontTable.xml",
        "word/numbering.xml",
        "word/settings.xml",
        "word/sectPr.xml",
    ):
        compared[part] = {
            "template": template_parts.get(part),
            "reference": reference_parts.get(part),
            "manuscript": manuscript_parts.get(part),
        }

    data = {
        "template": {
            "path": str(template),
            "sha256": sha256(template),
            "size": template.stat().st_size,
            "part_count": len(template_parts),
        },
        "reference": document_record(reference),
        "manuscript": document_record(manuscript),
        "part_comparison": compared,
        "template_parts": template_parts,
        "reference_parts": reference_parts,
    }
    output.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    print(json.dumps({
        "template_sha256": data["template"]["sha256"],
        "reference_sections": data["reference"]["sections"],
        "reference_styles": data["reference"]["styles"],
        "manuscript_sections": data["manuscript"]["sections"],
        "manuscript_tables": data["manuscript"]["tables"],
        "manuscript_inline_shapes": data["manuscript"]["inline_shape_count"],
        "manuscript_abstract_words": data["manuscript"]["abstract_words"],
        "manuscript_all_text_words": data["manuscript"]["all_text_words"],
        "manuscript_core_properties": data["manuscript"]["core_properties"],
        "styles_match": (
            data["part_comparison"]["word/styles.xml"]["template"]["sha256"]
            == data["part_comparison"]["word/styles.xml"]["manuscript"]["sha256"]
        ),
        "theme_match": (
            data["part_comparison"]["word/theme/theme1.xml"]["template"]["sha256"]
            == data["part_comparison"]["word/theme/theme1.xml"]["manuscript"]["sha256"]
        ),
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

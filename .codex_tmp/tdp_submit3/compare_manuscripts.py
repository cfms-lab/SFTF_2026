from __future__ import annotations

import difflib
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def paragraph_xml_text(paragraph) -> str:
    chunks: list[str] = []
    for node in paragraph._p.iter():
        if node.tag in {qn("w:t"), qn("m:t"), qn("w:tab"), qn("w:br")}:
            if node.tag == qn("w:tab"):
                chunks.append("\t")
            elif node.tag == qn("w:br"):
                chunks.append("\n")
            elif node.text:
                chunks.append(node.text)
    return "".join(chunks).strip()


def normalize(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_lines(path: Path) -> list[str]:
    doc = Document(path)
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = normalize(paragraph_xml_text(paragraph))
        if text:
            style = paragraph.style.name if paragraph.style else ""
            lines.append(f"[P:{style}] {text}")
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            cells = []
            for cell in row.cells:
                value = " / ".join(
                    normalize(paragraph_xml_text(p))
                    for p in cell.paragraphs
                    if normalize(paragraph_xml_text(p))
                )
                cells.append(value)
            lines.append(f"[T{ti}R{ri}] " + " || ".join(cells))
    return lines


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as package:
        return {
            name: hashlib.sha256(package.read(name)).hexdigest()
            for name in package.namelist()
            if name.startswith("word/media/")
        }


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("usage: compare_manuscripts.py TDP.docx RPJ.docx OUTDIR")
    tdp = Path(sys.argv[1])
    rpj = Path(sys.argv[2])
    outdir = Path(sys.argv[3])
    outdir.mkdir(parents=True, exist_ok=True)

    tdp_lines = extract_lines(tdp)
    rpj_lines = extract_lines(rpj)
    diff = list(
        difflib.unified_diff(
            tdp_lines,
            rpj_lines,
            fromfile=str(tdp),
            tofile=str(rpj),
            n=2,
            lineterm="",
        )
    )
    (outdir / "tdp_vs_rpj.diff.txt").write_text("\n".join(diff), encoding="utf-8")
    (outdir / "tdp_text.txt").write_text("\n".join(tdp_lines), encoding="utf-8")
    (outdir / "rpj_text.txt").write_text("\n".join(rpj_lines), encoding="utf-8")

    tdp_media = media_hashes(tdp)
    rpj_media = media_hashes(rpj)
    common_media_hashes = set(tdp_media.values()) & set(rpj_media.values())
    summary = {
        "tdp_line_count": len(tdp_lines),
        "rpj_line_count": len(rpj_lines),
        "diff_line_count": len(diff),
        "tdp_media_count": len(tdp_media),
        "rpj_media_count": len(rpj_media),
        "common_media_count": len(common_media_hashes),
        "tdp_only_media": [
            name for name, digest in tdp_media.items() if digest not in common_media_hashes
        ],
        "rpj_only_media": [
            name for name, digest in rpj_media.items() if digest not in common_media_hashes
        ],
    }
    (outdir / "comparison_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

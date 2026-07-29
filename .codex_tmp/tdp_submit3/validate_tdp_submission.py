from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from PIL import Image


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "m": M, "a": A, "wp": WP, "r": R, "pr": PR}


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def xml_text(element) -> str:
    return "".join(
        node.text or ""
        for node in element.xpath(".//*[local-name()='t']")
    )


def words(text: str) -> list[str]:
    return re.findall(r"[A-Za-z0-9]+(?:[–-][A-Za-z0-9]+)*", text)


def paragraph_texts(doc: Document) -> list[tuple[str, str]]:
    return [(p.style.name, xml_text(p._p)) for p in doc.paragraphs]


def expand_citation(token: str) -> set[int]:
    result: set[int] = set()
    for part in token.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, end = [int(x) for x in part.split("-", 1)]
            result.update(range(start, end + 1))
        else:
            result.add(int(part))
    return result


def main_doc_checks(path: Path, template: Path) -> dict:
    doc = Document(path)
    pars = paragraph_texts(doc)
    all_text = "\n".join(text for _, text in pars)
    idx = {text: i for i, (_, text) in enumerate(pars)}

    abstract_i = next(i for i, (_, text) in enumerate(pars) if text == "Abstract")
    keywords_i = next(i for i, (_, text) in enumerate(pars) if text.startswith("Keywords:"))
    abstract_text = " ".join(text for _, text in pars[abstract_i + 1 : keywords_i])
    intro_i = next(i for i, (_, text) in enumerate(pars) if text == "Introduction")
    acknowledgments_i = next(
        i for i, (_, text) in enumerate(pars) if text == "Acknowledgments"
    )
    main_count_text = " ".join(
        text
        for style, text in pars[intro_i:acknowledgments_i]
        if style not in {"Figure캡션", "Table캡션", "Figure Caption", "Table Caption"}
        and text
    )
    refs_i = next(i for i, (_, text) in enumerate(pars) if text == "References")
    body_before_refs = "\n".join(text for _, text in pars[:refs_i])
    ref_pars = [text for _, text in pars[refs_i + 1 :] if re.match(r"^\d+\. ", text)]
    figure_captions = [
        text for _, text in pars if re.match(r"^\ufeff?Figure \d+\.", text)
    ]
    table_captions = [
        text for _, text in pars if re.match(r"^\ufeff?Table \d+\.", text)
    ]
    keywords = [
        item.strip()
        for item in pars[keywords_i][1].split(":", 1)[1].split(";")
        if item.strip()
    ]

    required_headings = [
        "Abstract",
        "Introduction",
        "Materials and Methods",
        "Results",
        "Discussion",
        "Conclusions",
        "Acknowledgments",
        "Author Contributions",
        "Statements and Declarations",
        "Ethical considerations",
        "Consent to participate",
        "Consent for publication",
        "Declaration of conflicting interest",
        "Funding statement",
        "Data availability",
        "References",
    ]

    with ZipFile(path) as z:
        document_xml = z.read("word/document.xml")
        root = etree.fromstring(document_xml)
        superscript_tokens = []
        for run in root.xpath("//w:r[w:rPr/w:vertAlign[@w:val='superscript']]", namespaces=NS):
            value = "".join(run.xpath(".//w:t/text()", namespaces=NS))
            if re.fullmatch(r"\d+(?:-\d+)?(?:,\d+(?:-\d+)?)*", value):
                superscript_tokens.append(value)
        cited = set()
        for token in superscript_tokens:
            cited.update(expand_citation(token))

        math_paragraphs = len(root.xpath("//m:oMathPara", namespaces=NS))
        math_objects = len(root.xpath("//m:oMath", namespaces=NS))
        tracked_changes = len(root.xpath("//w:ins|//w:del", namespaces=NS))
        comments_present = "word/comments.xml" in z.namelist()
        image_names = [
            name
            for name in z.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        ]

        rel_root = etree.fromstring(z.read("word/_rels/document.xml.rels"))
        rels = {
            rel.get("Id"): rel.get("Target")
            for rel in rel_root.xpath("//pr:Relationship", namespaces=NS)
        }
        dpi_values = []
        for drawing in root.xpath("//w:drawing", namespaces=NS):
            blip = drawing.find(".//{%s}blip" % A)
            extent = drawing.find(".//{%s}extent" % WP)
            if blip is None or extent is None:
                continue
            rid = blip.get("{%s}embed" % R)
            target = rels.get(rid)
            if not target or not target.startswith("media/"):
                continue
            media_name = "word/" + target
            if media_name not in z.namelist():
                continue
            with Image.open(io.BytesIO(z.read(media_name))) as image:
                px_w, px_h = image.size
            width_in = int(extent.get("cx")) / 914400
            height_in = int(extent.get("cy")) / 914400
            dpi_values.append(
                {
                    "image": media_name,
                    "pixels": [px_w, px_h],
                    "display_inches": [round(width_in, 3), round(height_in, 3)],
                    "dpi": [round(px_w / width_in), round(px_h / height_in)],
                }
            )

        styles_hash = sha256(z.read("word/styles.xml"))
        theme_hash = sha256(z.read("word/theme/theme1.xml"))
        numbering_hash = sha256(z.read("word/numbering.xml"))

    with ZipFile(template) as z:
        template_styles_hash = sha256(z.read("word/styles.xml"))
        template_theme_hash = sha256(z.read("word/theme/theme1.xml"))
        template_numbering_hash = sha256(z.read("word/numbering.xml"))

    required_accessible_phrases = [
        "In practical terms, the normal tensor describes",
        "A useful way to read Eq. (2) is as a weighted vote.",
        "The virtual ground node is bookkeeping",
        "Uniform search spends the verification budget evenly",
    ]

    checks = {
        "abstract_words": len(words(abstract_text)),
        "abstract_under_300": len(words(abstract_text)) < 300,
        "main_text_words_approx": len(words(main_count_text)),
        "main_text_under_4000": len(words(main_count_text)) <= 4000,
        "keyword_count": len(keywords),
        "keywords_at_least_4": len(keywords) >= 4,
        "figure_caption_count": len(figure_captions),
        "figures_within_limit": len(figure_captions) <= 8,
        "table_caption_count": len(table_captions),
        "tables_within_limit": len(table_captions) <= 5,
        "reference_count": len(ref_pars),
        "references_within_limit": len(ref_pars) <= 100,
        "references_are_1_to_23": [int(re.match(r"^(\d+)\.", x).group(1)) for x in ref_pars]
        == list(range(1, 24)),
        "superscript_citation_groups": superscript_tokens,
        "cited_reference_numbers": sorted(cited),
        "citation_reference_1_to_1": cited == set(range(1, 24)),
        "author_year_citations_remaining": re.findall(
            r"\([A-ZÀ-ÖØ-Þ][^()]{0,150},\s*(?:19|20)\d{2}[a-z]?(?:;[^()]*)?\)",
            body_before_refs,
        ),
        "required_headings_missing": [h for h in required_headings if h not in idx],
        "coi_exact": (
            "The author(s) declared no potential conflicts of interest with respect "
            "to the research, authorship, and/or publication of this article"
        )
        in all_text,
        "funding_present": "NRF-2022R1A2C1010072" in all_text,
        "accurate_repository_wording": (
            "Access can be provided to the editor and reviewers during peer review"
            in all_text
            and "public repository" not in all_text
        ),
        "ai_disclosure_present": (
            "OpenAI Codex" in all_text
            and "Anthropic Claude Code" in all_text
            and "No generative AI was used to generate research data" in all_text
        ),
        "accessible_explanations_missing": [
            phrase for phrase in required_accessible_phrases if phrase not in all_text
        ],
        "strength_emphasis_present": (
            "The main practical advantage" in all_text
            and "roughly 29 times faster" in all_text
            and "targeted warm start" in all_text
        ),
        "figure_caption_numbers": [
            int(re.search(r"Figure (\d+)", x).group(1)) for x in figure_captions
        ],
        "table_caption_numbers": [
            int(re.search(r"Table (\d+)", x).group(1)) for x in table_captions
        ],
        "embedded_image_count": len(image_names),
        "display_math_paragraph_count": math_paragraphs,
        "math_object_count": math_objects,
        "tracked_changes_count": tracked_changes,
        "comments_part_present": comments_present,
        "image_dpi": dpi_values,
        "minimum_effective_dpi": min(
            min(item["dpi"]) for item in dpi_values
        )
        if dpi_values
        else None,
        "styles_match_template": styles_hash == template_styles_hash,
        "theme_matches_template": theme_hash == template_theme_hash,
        "numbering_matches_template": numbering_hash == template_numbering_hash,
        "core_properties": {
            "title": doc.core_properties.title,
            "author": doc.core_properties.author,
            "subject": doc.core_properties.subject,
        },
    }
    return checks


def cover_checks(path: Path) -> dict:
    doc = Document(path)
    text = "\n".join(xml_text(p._p) for p in doc.paragraphs)
    return {
        "paragraph_count": len(doc.paragraphs),
        "date_present": "28 July 2026" in text,
        "journal_present": "3D Printing and Additive Manufacturing" in text,
        "article_type_present": "Original Article" in text,
        "title_present": "The Support Flow Tensor Field" in text,
        "novelty_and_fit_present": (
            "4.57 s per mesh" in text
            and "60 independently selected" in text
            and "complementary problem" in text
        ),
        "originality_statement_present": (
            "has not been published" in text
            and "not under consideration elsewhere" in text
        ),
        "ai_disclosure_present": (
            "OpenAI Codex" in text
            and "Anthropic Claude Code" in text
            and "also disclosed in the manuscript" in text
        ),
        "coi_present": "declares no potential conflicts of interest" in text,
        "funding_present": "NRF-2022R1A2C1010072" in text,
        "reviewer_recommendation_absent": "recommended reviewer" not in text.lower(),
        "author_contact_present": (
            "snowman0@kumoh.ac.kr" in text
            and "0000-0003-0105-920X" in text
        ),
        "core_properties": {
            "title": doc.core_properties.title,
            "author": doc.core_properties.author,
            "subject": doc.core_properties.subject,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--main", type=Path, required=True)
    parser.add_argument("--cover", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    report = {
        "main": main_doc_checks(args.main, args.template),
        "cover": cover_checks(args.cover),
    }
    args.out.write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
